"""
Tests for GET /export_xlsx (app/routes/reports.py) -- specifically the
per-player "Clean rounds" figure added to the Summary sheet.
"""

import io

import openpyxl
import pytest

from engine.referee import RefereeSession
from tests.conftest import make_player
from app import create_app
from app.models.game_room import GameRoom, GameConfig
from app.services.session_store import game_sessions, set_session


@pytest.fixture
def app():
    return create_app()


@pytest.fixture
def client(app):
    return app.test_client()


def _make_room_with_export_data(room_code="TestExport"):
    players = [make_player("Alice"), make_player("Bob")]
    players[0].is_dealer = True
    raw = RefereeSession(players, "Alice", wager=1, num_hands=1)
    room = GameRoom(session=raw, config=GameConfig(mode="digital", drinking_mode=True))

    # Two rounds' worth of drink-log rows so num_rounds == 2. Alice needs at
    # least one row too (players_seen is built from csv_rows) even though
    # she stayed clean -- give her an Ace-of-Clubs style credit row instead
    # of a plain drink, so she still shows up with 0 net sips overall.
    room.drinks.csv_rows = [
        {"round": 1, "dealer": "Alice", "player": "Bob", "role": "player",
         "rule": "Net hand losses", "sips": 2},
        {"round": 2, "dealer": "Alice", "player": "Bob", "role": "player",
         "rule": "Net hand losses", "sips": 1},
        {"round": 1, "dealer": "Alice", "player": "Alice", "role": "player",
         "rule": "Ace of Clubs credit", "sips": 0},
    ]
    room.stats.total_clean_rounds = {"Alice": 2, "Bob": 0}   # Alice clean both rounds, Bob neither

    set_session(room_code, room)
    return room_code, room


def _summary_rows(resp_data):
    wb = openpyxl.load_workbook(io.BytesIO(resp_data))
    ws = wb["Summary"]
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_export_includes_clean_rounds_for_each_player(client):
    room_code, room = _make_room_with_export_data()
    try:
        resp = client.get(f"/export_xlsx?room_code={room_code}")
        assert resp.status_code == 200

        rows = _summary_rows(resp.data)
        joined = [" | ".join(str(c) for c in row if c is not None) for row in rows]

        assert any("Clean rounds: 2/2 (100.0%)" in line for line in joined)   # Alice
        assert any("Clean rounds: 0/2 (0.0%)" in line for line in joined)     # Bob
    finally:
        game_sessions.pop(room_code, None)


def test_export_clean_rounds_defaults_to_zero_when_untracked(client):
    """A player with no entry in total_clean_rounds (e.g. never had a
    harvested round) must not KeyError -- shows as 0, not crash."""
    room_code, room = _make_room_with_export_data()
    room.stats.total_clean_rounds = {}   # nobody tracked
    try:
        resp = client.get(f"/export_xlsx?room_code={room_code}")
        assert resp.status_code == 200
        rows = _summary_rows(resp.data)
        joined = [" | ".join(str(c) for c in row if c is not None) for row in rows]
        assert any("Clean rounds: 0/2 (0.0%)" in line for line in joined)
    finally:
        game_sessions.pop(room_code, None)


def test_export_includes_raw_sips_and_drink_rate_section(client):
    room_code, room = _make_room_with_export_data()
    room.drinks.csv_rows = [
        {"round": 1, "dealer": "Alice", "player": "Bob", "role": "player",
         "rule": "Net hand losses", "sips": 2, "sips_raw": 3, "rate": 0.6},
        {"round": 2, "dealer": "Alice", "player": "Bob", "role": "player",
         "rule": "Net hand losses", "sips": 1, "sips_raw": 1, "rate": 1.0},
    ]
    room.stats.rate_rounds = {"Bob": {0.6: 1, 1.0: 1}, "Alice": {1.0: 2}}
    try:
        resp = client.get(f"/export_xlsx?room_code={room_code}")
        rows = _summary_rows(resp.data)
        joined = [" | ".join(str(c) for c in row if c is not None) for row in rows]
        assert any("Raw sips: 4" in line and "Avg drink factor: 0.80×" in line for line in joined)
        assert any(line.startswith("DRINK RATE") for line in joined)
        assert any(line.startswith("Bob | 2 | 0.8 | 4 | 3 | 0.75") for line in joined)
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        header = [c.value for c in wb["Raw Rows"][1]]
        assert header[-2:] == ["sips_raw", "rate"]
    finally:
        game_sessions.pop(room_code, None)
