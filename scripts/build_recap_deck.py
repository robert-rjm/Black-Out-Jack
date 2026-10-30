"""
scripts/build_recap_deck.py
============================
Build the session recap deck (app/services/recap_deck.py) as a vector PDF
from a session's two saved exports (drinks_summary_<day>.xlsx +
decision_log_<day>.xlsx in data/decisions/), printed by a local headless
Edge/Chrome. The in-app button (GET /recap_html) renders the same deck in
the browser instead.

Usage:
    python scripts/build_recap_deck.py 2026-10-19
    python scripts/build_recap_deck.py 2026-10-19 --out data/recaps --html
"""

import argparse
import os
import shutil
import subprocess
import sys
import tempfile
import time

import openpyxl

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)

from app.services.recap_deck import analyse, build_html, logo_css  # noqa: E402


def load(day: str, folder: str) -> dict:
    drinks = openpyxl.load_workbook(os.path.join(folder, f"drinks_summary_{day}.xlsx"), read_only=True)
    raw = [r for r in drinks["Raw Rows"].iter_rows(min_row=2, values_only=True) if r[0]]
    rows = [dict(round=r[0], dealer=r[1], player=r[2], role=r[3], rule=r[4], sips=r[5]) for r in raw]

    dec_wb = openpyxl.load_workbook(os.path.join(folder, f"decision_log_{day}.xlsx"), read_only=True)
    ws = dec_wb["Hand Decisions"]
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    decisions = [dict(zip(hdr, r)) for r in it if r[0]]
    ws = dec_wb["Dealer Lottery Entries"]
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    lottery = [dict(zip(hdr, r)) for r in it if r[0]]
    return analyse(rows, decisions, lottery, day)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _browser() -> str | None:
    for path in (r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
                 r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
                 r"C:\Program Files\Google\Chrome\Application\chrome.exe"):
        if os.path.exists(path):
            return path
    return shutil.which("chromium") or shutil.which("google-chrome") or shutil.which("msedge")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("day", help="session date as in the export file names, e.g. 2026-10-19")
    ap.add_argument("--dir", default=os.path.join(_ROOT, "data", "decisions"))
    ap.add_argument("--out", default=os.path.join(_ROOT, "data", "recaps"))
    ap.add_argument("--html", action="store_true", help="also keep the HTML next to the PDF")
    args = ap.parse_args()

    a = load(args.day, args.dir)
    os.makedirs(args.out, exist_ok=True)
    html_path = os.path.join(args.out if args.html else tempfile.mkdtemp(), f"recap_{args.day}.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(build_html(a, logo_css(inline=True)))
    pdf_path = os.path.join(args.out, f"recap_{args.day}.pdf")
    exe = _browser()
    if not exe:
        sys.exit(f"No Edge/Chrome found; HTML written to {html_path}")
    if os.path.exists(pdf_path):
        os.remove(pdf_path)
    # A throwaway profile, or an already-open Edge/Chrome swallows the call;
    # the launcher can exit before the PDF lands, so wait for the file
    profile = tempfile.mkdtemp()
    subprocess.run([exe, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--user-data-dir={profile}", f"--print-to-pdf={pdf_path}",
                    "file:///" + html_path.replace("\\", "/")],
                   capture_output=True, timeout=120)
    for _ in range(60):
        if os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 0:
            break
        time.sleep(0.5)
    shutil.rmtree(profile, ignore_errors=True)
    if not os.path.exists(pdf_path):
        sys.exit(f"PDF printing failed; HTML written to {html_path}")
    print(f"Wrote {pdf_path}")


if __name__ == "__main__":
    main()
